"""Build synchronized Word and LaTeX from manuscript/content.json.

Word equations are rendered from the same LaTeX expressions used in paper.tex.
Compile paper.tex with pdflatex twice for the designated paper.pdf.
"""
from pathlib import Path
import json, re, shutil
from datetime import datetime, timezone
from xml.sax.saxutils import escape
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
M=ROOT/'manuscript'

def tex(s):
    replacements={'\\':r'\textbackslash{}','&':r'\&','%':r'\%','$':r'\$','#':r'\#','_':r'\_','{':r'\{','}':r'\}','~':r'\textasciitilde{}','^':r'\textasciicircum{}'}
    replacements.update({'η':r'\ensuremath{\eta}','τ':r'\ensuremath{\tau}','ρ':r'\ensuremath{\rho}','π':r'\ensuremath{\pi}','Ω':r'\ensuremath{\Omega}','∂':r'\ensuremath{\partial}','∞':r'\ensuremath{\infty}','×':r'\ensuremath{\times}','≤':r'\ensuremath{\leq}','≥':r'\ensuremath{\geq}','µ':r'\ensuremath{\mu}','′':r'\ensuremath{\prime}','−':r'\ensuremath{-}','–':'--','—':'---'})
    for c,v in zip('⁰¹²³⁴⁵⁶⁷⁸⁹⁻','0123456789-'):replacements[c]=r'\textsuperscript{'+v+'}'
    for c,v in zip('₀₁₂₃₄₅₆₇₈₉','0123456789'):replacements[c]=r'\textsubscript{'+v+'}'
    return ''.join(replacements.get(c,c) for c in s)

def build():
    content=json.loads((M/'content.json').read_text(encoding='utf8'))
    assets=M/'assets';assets.mkdir(exist_ok=True)
    doc=Document(); sec=doc.sections[0]
    sec.top_margin=sec.bottom_margin=Inches(.8)
    sec.left_margin=sec.right_margin=Inches(.85)
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    for st in doc.styles:
        for el in list(st.element.iter(qn('w:pBdr'))):
            el.getparent().remove(el)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
        st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        rf=st.element.get_or_add_rPr().rFonts
        if rf is not None:
            for a in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
                rf.attrib.pop(qn('w:'+a),None)
    normal=doc.styles['Normal'];normal.font.size=Pt(11)
    normal.paragraph_format.line_spacing=1.08
    normal.paragraph_format.space_after=Pt(6)
    doc.styles['Title'].font.size=Pt(20)
    doc.styles['Heading 1'].font.size=Pt(14)
    doc.styles['Heading 2'].font.size=Pt(12)
    for name in ['Heading 1','Heading 2']:
        doc.styles[name].paragraph_format.keep_with_next=True
    doc.core_properties.author='UCP Technology LLC'
    doc.core_properties.created=datetime.now(timezone.utc)
    doc.core_properties.modified=datetime.now(timezone.utc)
    doc.core_properties.title=content['title']
    doc.core_properties.subject='Conditional insulin transport calculation and dose-preserving geometry frontier'
    doc.core_properties.keywords='islet encapsulation; insulin transport; diffusion; human scale'
    doc.add_paragraph(content['title'],'Title')
    doc.add_paragraph('UCP Technology LLC')
    doc.add_paragraph('26 October 2026')
    footer=sec.footer.paragraphs[0];footer.alignment=1
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    latex=[r'\documentclass[11pt]{article}',r'\usepackage[letterpaper,margin=0.85in]{geometry}',
      r'\usepackage[T1]{fontenc}',r'\usepackage{lmodern}',r'\usepackage{amsmath,amssymb,graphicx,booktabs,array,tabularx}',
      r'\usepackage[hidelinks]{hyperref}',r'\usepackage{microtype}',r'\setlength{\parindent}{0pt}',
      r'\setlength{\parskip}{4pt}',r'\setlength{\textfloatsep}{10pt}\setlength{\intextsep}{10pt}',r'\makeatletter\renewcommand\section{\@startsection{section}{1}{0pt}{-2ex plus -.5ex minus -.2ex}{1ex plus .2ex}{\normalfont\Large\bfseries}}\makeatother',r'\setlength{\emergencystretch}{2em}',r'\renewcommand{\topfraction}{0.9}\renewcommand{\textfraction}{0.08}',r'\hypersetup{pdftitle={'+tex(content['title'])+r'},pdfauthor={UCP Technology LLC}}',
      r'\begin{document}',r'\begin{center}{\LARGE\bfseries '+tex(content['title'])+r'}\\[10pt]UCP Technology LLC\\26 October 2026\end{center}']
    eq_index=0;fig_index=0;end_matter=False
    for b in content['blocks']:
        kind=b['type']
        if kind=='heading':
            if b['text']=='10 Conclusion':
                latex.append(r'\par\noindent\begin{minipage}{\linewidth}');end_matter=True
            doc.add_paragraph(b['text'],'Heading 1')
            latex.append(r'\section*{'+tex(b['text'])+'}')
        elif kind=='paragraph':
            doc.add_paragraph(b['text'])
            latex.append(tex(b['text'])+'\n')
        elif kind=='equation':
            eq_index+=1;name=f'equation_{eq_index:02d}.png'
            fig=plt.figure(figsize=(7,.6));fig.text(.5,.5,'$'+b['latex']+'$',ha='center',va='center',fontsize=13)
            fig.savefig(assets/name,dpi=220,bbox_inches='tight',pad_inches=.04,facecolor='white');plt.close(fig)
            p=doc.add_paragraph();p.alignment=1
            from PIL import Image
            with Image.open(assets/name) as im: width=min(6.0,im.width/220)
            p.add_run().add_picture(str(assets/name),width=Inches(width))
            latex.append(r'\begin{equation}'+b['latex']+r'\end{equation}')
            if b.get('explanation'):
                doc.add_paragraph(b['explanation']);latex.append(tex(b['explanation']))
        elif kind=='figure':
            fig_index+=1
            source=ROOT/b['file'];target=assets/source.name;shutil.copy2(source,target)
            p=doc.add_paragraph();p.alignment=1;p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(str(target),width=Inches(b.get('width_in',6.3)))
            cap=doc.add_paragraph(f'Figure {fig_index}. '+b['caption'],'Caption');cap.paragraph_format.space_after=Pt(10)
            latex.append(r'\begin{figure}[htbp]\centering\includegraphics[width='+str(min(1,b.get('width_in',6.3)/6.8))+r'\linewidth]{assets/'+source.name+r'}\caption{'+tex(b['caption'])+r'}\end{figure}')
        elif kind=='table':
            cap=doc.add_paragraph(b['caption'],'Caption');cap.paragraph_format.keep_with_next=True
            table=doc.add_table(rows=1,cols=len(b['headers']));table.style='Table Grid'
            borders=OxmlElement('w:tblBorders')
            for edge in ['top','left','bottom','right','insideH','insideV']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            table._tbl.tblPr.append(borders)
            for cell,text in zip(table.rows[0].cells,b['headers']):
                cell.text=text
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDEDED');cell._tc.get_or_add_tcPr().append(sh)
                for run in cell.paragraphs[0].runs:run.font.bold=True;run.font.color.rgb=RGBColor(0,0,0)
            for row in b['rows']:
                for cell,text in zip(table.add_row().cells,row):cell.text=str(text)
            for row in table.rows:
                row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after=Pt(3)
                        for r in p.runs:r.font.size=Pt(10)
            props=table.rows[0]._tr.get_or_add_trPr();props.append(OxmlElement('w:tblHeader'))
            latex.append(r'\par\begin{table}[htbp]\centering\small '+tex(b['caption'])+r'\\[4pt]\begin{tabularx}{\linewidth}{'+r'>{\raggedright\arraybackslash}X'*len(b['headers'])+r'}\toprule')
            latex.append(' & '.join(tex(x) for x in b['headers'])+r'\\\midrule')
            for row in b['rows']:latex.append(' & '.join(tex(str(x)) for x in row)+r'\\')
            latex.append(r'\bottomrule\end{tabularx}\end{table}\par')
        elif kind=='reference':
            p=doc.add_paragraph(b['text']);p.paragraph_format.space_after=Pt(3)
            p.paragraph_format.keep_together=True
            for run in p.runs:run.font.size=Pt(9)
            latex.append(r'\par\noindent\begin{minipage}{\linewidth}\small '+tex(b['text'])+r'\end{minipage}\par\vspace{2pt}'+'\n')
    if end_matter:latex.append(r'\end{minipage}')
    latex.append(r'\end{document}')
    doc.save(M/'paper.docx')
    (M/'paper.tex').write_text('\n'.join(latex)+'\n',encoding='utf8')
    print('Built synchronized paper.docx and paper.tex from content.json')

if __name__=='__main__':build()
